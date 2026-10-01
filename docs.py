"""Document intelligence: text extraction, statistics and keyword search."""
import re
from collections import Counter
from markupsafe import Markup, escape

STOP = set("""a about above after again all also am an and any are as at be because been before being below between both but by can
could did do does doing down during each few for from further had has have having he her here hers him his how i if in into is it its
just me more most my no nor not now of off on once only or other our out over own same she should so some such than that the their
them then there these they this those through to too under until up very was we were what when where which while who whom why will
with would you your""".split())


class DocError(Exception):
    pass


def extract_text(path, ext):
    pages = None
    try:
        if ext == "txt":
            raw = open(path, "rb").read()
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
        elif ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            pages = len(reader.pages)
            text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        elif ext == "docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n\n".join(parts)
        else:
            raise DocError("Unsupported document type. Use PDF, DOCX or TXT.")
    except DocError:
        raise
    except Exception as e:
        raise DocError(f"Could not read this document: {str(e)[:120]}")
    text = text.replace("\x00", "").strip()
    if not text:
        raise DocError("No text could be extracted (the file may be empty or a scanned image).")
    return text, pages


def keywords(text, n=15):
    words = [w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'-]{2,}", text)]
    return Counter(w for w in words if w not in STOP).most_common(n)


def stats(text, pages=None):
    words = re.findall(r"[A-Za-z0-9'-]+", text)
    sentences = [s for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    nw = len(words)
    return dict(words=nw, chars=len(text), sentences=len(sentences), paragraphs=len(paragraphs), pages=pages,
                unique_words=len({w.lower() for w in words}),
                avg_word=round(sum(len(w) for w in words) / nw, 1) if nw else 0,
                avg_sentence=round(nw / len(sentences), 1) if sentences else 0,
                reading=max(1, round(nw / 200)), keywords=keywords(text))


def search(text, query, width=70, limit=40):
    query = (query or "").strip()
    if not query:
        return 0, []
    hits = list(re.finditer(re.escape(query), text, re.I))
    out = []
    for m in hits[:limit]:
        a, b = max(0, m.start() - width), min(len(text), m.end() + width)
        snippet = (("…" if a else "") + escape(text[a:m.start()]) + Markup("<mark>") + escape(text[m.start():m.end()])
                   + Markup("</mark>") + escape(text[m.end():b]) + ("…" if b < len(text) else ""))
        out.append(Markup(snippet.replace("\n", " ")))
    return len(hits), out
